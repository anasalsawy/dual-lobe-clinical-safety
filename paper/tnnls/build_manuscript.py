"""Build the TNNLS manuscript (.docx) from the official IEEE Transactions Word template.

The template's styles, page geometry (two columns, 0.65 in margins), header,
first-page footnote and reference numbering are kept; only the content is
replaced. Run `python make_figures.py` first, then this script. Bracketed
text highlighted in yellow is information only the author can supply.
"""
import copy
import re
from pathlib import Path
from xml.sax.saxutils import escape

import docx
from docx.shared import Inches
from lxml import etree

HERE = Path(__file__).parent
TEMPLATE = HERE / "template" / "IEEE-Transactions-template.docx"
OUT = HERE / "TNNLS_manuscript.docx"
FIG = HERE / "figures"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = f'xmlns:w="{W}" xmlns:m="{M}"'

AUTHOR = "Anas Alsawy"
TITLE = "Surmounting the Open Sesame Barrier in Healthcare Applications"
RUNNING_HEAD = "ALSAWY: SURMOUNTING THE OPEN SESAME BARRIER IN HEALTHCARE APPLICATIONS"


# --------------------------------------------------------------------------
# Inline markup:  *italic*   _{subscript}   ^{superscript}   [[placeholder]]
# --------------------------------------------------------------------------
TOKEN = re.compile(r"(\*[^*]+\*|_i?\{[^}]*\}|\^i?\{[^}]*\}|\[\[[^\]]+\]\])")


def run(text, italic=False, bold=False, sub=False, sup=False, size=None, smallcaps=False,
        highlight=False):
    rpr = ""
    if bold:
        rpr += "<w:b/>"
    if italic:
        rpr += "<w:i/>"
    if smallcaps:
        rpr += "<w:smallCaps/>"
    if highlight:
        rpr += '<w:highlight w:val="yellow"/>'
    if size:
        rpr += f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>'
    if sub:
        rpr += '<w:vertAlign w:val="subscript"/>'
    if sup:
        rpr += '<w:vertAlign w:val="superscript"/>'
    text = text.replace("'", "\u2019")
    return (f'<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r>')


def runs(text, bold=False, size=None, base_italic=False):
    out = []
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("*"):
            out.append(run(tok[1:-1], italic=not base_italic, bold=bold, size=size))
        elif tok.startswith("_i{"):
            out.append(run(tok[3:-1], sub=True, bold=bold, size=size, italic=True))
        elif tok.startswith("^i{"):
            out.append(run(tok[3:-1], sup=True, bold=bold, size=size, italic=True))
        elif tok.startswith("_{"):
            out.append(run(tok[2:-1], sub=True, bold=bold, size=size, italic=base_italic))
        elif tok.startswith("^{"):
            out.append(run(tok[2:-1], sup=True, bold=bold, size=size, italic=base_italic))
        elif tok.startswith("[["):
            out.append(run("[" + tok[2:-2] + "]", bold=bold, size=size, highlight=True))
        else:
            out.append(run(tok, bold=bold, size=size, italic=base_italic))
    return "".join(out)


def el(xml):
    return etree.fromstring(f"<root {NS}>{xml}</root>")[0]


BODY_PPR = ('<w:pPr><w:pStyle w:val="Text"/><w:widowControl w:val="0"/>'
            '<w:spacing w:line="252" w:lineRule="auto"/><w:ind w:firstLine="202"/>'
            '<w:jc w:val="both"/></w:pPr>')


def P(text):
    return el(f"<w:p>{BODY_PPR}{runs(text)}</w:p>")


NOINDENT_PPR = BODY_PPR.replace('w:firstLine="202"', 'w:firstLine="0"')


def P_noindent(text):
    return el("<w:p>" + NOINDENT_PPR + runs(text) + "</w:p>")


def H1(num, text):
    return el(f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>{run(f"{num}. {text}")}</w:p>')


def H2(letter, text):
    return el('<w:p><w:pPr><w:pStyle w:val="Heading2"/><w:numPr><w:ilvl w:val="0"/>'
              f'<w:numId w:val="0"/></w:numPr></w:pPr>{run(f"{letter}. {text}")}</w:p>')


def ITEM(label, text):
    """Hanging-indent list item, e.g. '1)' or a bullet."""
    return el('<w:p><w:pPr><w:pStyle w:val="Text"/><w:widowControl w:val="0"/>'
              '<w:tabs><w:tab w:val="left" w:pos="562"/></w:tabs>'
              '<w:spacing w:line="252" w:lineRule="auto"/>'
              '<w:ind w:left="562" w:hanging="360"/><w:jc w:val="both"/></w:pPr>'
              f'{run(label)}<w:r><w:tab/></w:r>{runs(text)}</w:p>')


def DROPCAP(first_word, rest):
    """IEEE drop-cap opening paragraph (the template's framed large initial)."""
    cap = ('<w:p><w:pPr><w:keepNext/><w:framePr w:dropCap="drop" w:lines="2" w:wrap="around" '
           'w:vAnchor="text" w:hAnchor="text"/><w:spacing w:line="481" w:lineRule="exact"/>'
           '<w:jc w:val="both"/><w:textAlignment w:val="baseline"/><w:rPr><w:position w:val="-4"/>'
           '<w:sz w:val="61"/></w:rPr></w:pPr><w:r><w:rPr><w:position w:val="-4"/><w:sz w:val="61"/>'
           f'</w:rPr><w:t>{first_word[0]}</w:t></w:r></w:p>')
    body = "<w:p>" + NOINDENT_PPR + run(first_word[1:].upper()) + runs(rest) + "</w:p>"
    return [el(cap), el(body)]


# ------------------------------ equations ---------------------------------
def EQ(text, number):
    """Display equation: centered via tab stops, number flush right (template style)."""
    return el('<w:p><w:pPr><w:widowControl w:val="0"/><w:tabs><w:tab w:val="center" w:pos="2520"/>'
              '<w:tab w:val="right" w:pos="5040"/></w:tabs><w:spacing w:before="60" w:after="60" '
              'w:line="252" w:lineRule="auto"/></w:pPr><w:r><w:tab/></w:r>'
              f'{runs(text)}<w:r><w:tab/><w:t>({number})</w:t></w:r></w:p>')


# ------------------------------ figures -----------------------------------
FIG_PARAS = []


def FIGURE(doc, filename, n, caption):
    p = doc.add_paragraph()
    p.alignment = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = docx.shared.Pt(6)
    p.add_run().add_picture(str(FIG / filename), width=Inches(3.4))
    elm = p._p
    elm.getparent().remove(elm)
    cap = el('<w:p><w:pPr><w:pStyle w:val="FigureCaption"/><w:spacing w:before="60" w:after="120"/>'
             f'</w:pPr>{run(f"Fig. {n}.", size=16)}{run("  ", size=16)}{runs(caption, size=16)}</w:p>')
    return [elm, cap]


# ------------------------------ references --------------------------------
REFS = [
    'A. Vaswani *et al.*, "Attention is all you need," in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 30, 2017, pp. 5998–6008.',
    'L. Ouyang *et al.*, "Training language models to follow instructions with human feedback," in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 35, 2022, pp. 27730–27744.',
    'O. Press, M. Zhang, S. Min, L. Schmidt, N. A. Smith, and M. Lewis, "Measuring and narrowing the compositionality gap in language models," in *Findings Assoc. Comput. Linguistics: EMNLP*, 2023, pp. 5687–5711.',
    'S. Yao *et al.*, "Tree of thoughts: Deliberate problem solving with large language models," in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 36, 2023.',
    'D. Zhou *et al.*, "Least-to-most prompting enables complex reasoning in large language models," in *Proc. Int. Conf. Learn. Represent. (ICLR)*, 2023.',
    'L. Wang *et al.*, "Plan-and-solve prompting: Improving zero-shot chain-of-thought reasoning by large language models," in *Proc. 61st Annu. Meeting Assoc. Comput. Linguistics (ACL)*, 2023, pp. 2609–2634.',
    'X. Ning, Z. Lin, Z. Zhou, Z. Wang, H. Yang, and Y. Wang, "Skeleton-of-thought: Prompting LLMs for efficient parallel generation," in *Proc. Int. Conf. Learn. Represent. (ICLR)*, 2024.',
    'E. Zelikman, G. Harik, Y. Shao, V. Jayasiri, N. Haber, and N. D. Goodman, "Quiet-STaR: Language models can teach themselves to think before speaking," in *Proc. Conf. Lang. Model. (COLM)*, 2024.',
    'J. Zhou, H. Li, S. Chen, Z. Chen, Z. Han, and X. Gao, "Large language models in biomedicine and healthcare," *npj Artif. Intell.*, vol. 1, 2025, Art. no. 44, doi: 10.1038/s44387-025-00047-1.',
    'Y. Liu *et al.*, "Benchmarking large language model-based agent systems for clinical decision tasks," *npj Digit. Med.*, 2026, doi: 10.1038/s41746-026-02443-6.',
    'J. Wei *et al.*, "Chain-of-thought prompting elicits reasoning in large language models," in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 35, 2022, pp. 24824–24837.',
    'T. Kojima, S. S. Gu, M. Reid, Y. Matsuo, and Y. Iwasawa, "Large language models are zero-shot reasoners," in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 35, 2022, pp. 22199–22213.',
    'A. Madaan *et al.*, "Self-refine: Iterative refinement with self-feedback," in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, vol. 36, 2023.',
    'S. Dhuliawala *et al.*, "Chain-of-verification reduces hallucination in large language models," in *Findings Assoc. Comput. Linguistics: ACL*, 2024, pp. 3563–3578.',
    'Y. Du, S. Li, A. Torralba, J. B. Tenenbaum, and I. Mordatch, "Improving factuality and reasoning in language models through multiagent debate," in *Proc. 41st Int. Conf. Mach. Learn. (ICML)*, 2024.',
    'Y. Bai *et al.*, "Constitutional AI: Harmlessness from AI feedback," 2022, *arXiv:2212.08073*.',
]


def REF(text):
    text = re.sub(r'"([^"]*)"', "\u201c\\1\u201d", text)
    return el('<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
              '<w:ind w:left="360" w:hanging="360"/><w:jc w:val="both"/><w:rPr><w:sz w:val="16"/>'
              f'<w:szCs w:val="16"/></w:rPr></w:pPr>{runs(text, size=16)}</w:p>')


def REF_HEAD():
    return el('<w:p><w:pPr><w:keepNext/><w:spacing w:before="240" w:after="80"/><w:jc w:val="center"/>'
              f'</w:pPr>{run("References", smallcaps=True)}</w:p>')


# ------------------------------ content -----------------------------------
ABSTRACT = (
    "Large language models (LLMs) are highly capable at following instructions, using tools, "
    "writing code, and carrying out multistep reasoning, yet their usefulness depends heavily on "
    "how precisely users phrase their requests and on the conversational groundwork users lay. "
    "This paper introduces the Open Sesame barrier: a behavioral constraint in which "
    "decision-critical information that the model can supply remains inactive until a user asks "
    "the question that elicits it. Models that excel on narrowly specified tasks can therefore "
    "fail in broader contexts, a serious problem in autonomous and clinical settings, where "
    "failing to raise an implicit but critical question can outweigh the accuracy of everything "
    "else that is said. We propose two training objectives: prompt-space prediction, which teaches "
    "a model to anticipate a structured set of latent follow-up questions, and coarse-to-fine "
    "context construction, which builds a hierarchical representation of the problem before "
    "detailed reasoning begins. Because greater autonomy also raises the cost of error, we further "
    "propose a dual-lobe supervisory architecture: a generator/executor lobe performs the "
    "reasoning, and an independent supervisory lobe searches for unaddressed questions, verifies "
    "claims, detects shifts in intent, and enforces domain safety rules. We conclude with testable "
    "hypotheses and an evaluation framework spanning software engineering, open-domain "
    "decision-making, and clinical decision support."
)
INDEX_TERMS = ("AI safety, clinical decision support, hierarchical planning, large language models "
               "(LLMs), medical AI, prompt prediction, question generation, reasoning, verification")


def body(doc):
    B = []
    add = B.append
    ext = B.extend

    # I. Introduction ------------------------------------------------------
    add(H1("I", "Introduction"))
    ext(DROPCAP("Modern", " artificial intelligence (AI) systems primarily operate through a simple "
        "interface: the user writes a prompt and the system returns a response. Behind this "
        "simplicity, contemporary large language models (LLMs) can tackle intricate tasks, including "
        "representation learning, inference, tool selection, planning, and multistep reasoning. "
        "Nevertheless, the quality of their output remains tightly coupled to the user's prompts and "
        "to the dialogue those prompts produce."))
    add(P("This coupling exposes an important limitation. Consider a user who collaborates "
          "extensively with an AI system to engineer a software feature, relying on the model for "
          "coding, debugging, comparing frameworks, and iteratively refining the design. Eventually "
          "the user discovers that a mature solution already exists. Once the user asks, “Is "
          "there an existing solution?” the model readily searches for and evaluates the "
          "relevant options."))
    add(P("The key observation is not that the model possessed this information the way a human "
          "expert would, but that the relevant information remained dormant until the user's "
          "question activated it. We call this phenomenon the *Open Sesame barrier*: like a locked "
          "door, crucial information becomes accessible only when the right key is presented. In "
          "conversational AI, that key is often a question the user does not realize they need to "
          "ask."))
    add(P("This reveals a fundamental mismatch. Users typically consult AI systems because they "
          "lack a complete picture of their situation, yet the success of the interaction depends "
          "on the user already knowing which omitted elements need clarification. As AI systems "
          "become more autonomous, this mismatch grows: stronger local competence can prolong an "
          "effort that is misdirected at the global level."))
    add(P("This paper makes four contributions:"))
    add(ITEM("1)", "We characterize the Open Sesame barrier as a prompt-dependent salience failure "
             "that coexists with strong local proficiency and complicates supervision."))
    add(ITEM("2)", "We propose *prompt-space prediction*, a training objective that teaches models "
             "to anticipate the latent questions associated with a user's intent before responding "
             "or acting."))
    add(ITEM("3)", "We introduce *coarse-to-fine context construction*, which guides models to build "
             "a hierarchical representation of the problem before committing reasoning effort."))
    add(ITEM("4)", "We present a *dual-lobe supervisory architecture* as a safety mechanism for "
             "systems that perform extended reasoning, with particular attention to clinical "
             "settings."))
    add(P("Our central claim is that the most effective remedy is neither better prompting nor "
          "additional agents alone. Rather, the focus should be on the training objectives and the "
          "behaviors they instill, with dual-lobe supervision serving as the accompanying safety "
          "architecture. The remainder of the paper is organized as follows. Section II develops "
          "the background, Section III reviews related work, Sections IV and V present the two "
          "training objectives, Section VI states the launchpad hypothesis, Section VII motivates "
          "medicine as a testbed, Section VIII describes dual-lobe supervision, Sections IX and X "
          "present the evaluation framework and hypotheses, and Sections XI–XIII discuss "
          "limitations and conclude."))

    # II. Background -------------------------------------------------------
    add(H1("II", "Background: From Conditional Continuation to Behavioral Limitation"))
    add(H2("A", "Continuation-Centric Training"))
    add(P("Autoregressive LLMs are trained to model the conditional probability of each token given "
          "the preceding tokens. Transformer architectures represent sequences effectively [1], and "
          "instruction tuning and reinforcement learning from human feedback (RLHF) [2] optimize "
          "them to produce outputs that satisfy explicit human requests. This methodology has "
          "yielded highly capable assistants."))
    add(P("We contend, however, that the underlying behavior remains continuation-centric: the "
          "provided context steers the model toward the expected next output rather than toward a "
          "complete account of the problem."))
    add(P("This design creates an imbalance between depth and breadth. Once a reasoning path is "
          "engaged, the model can reason thoroughly along it. Information outside that illuminated "
          "path, however, often stays dormant, waiting to be invoked by a later question, a "
          "retrieval call, or a user prompt designed to elicit it."))
    add(H2("B", "The Prompt-Echo Effect"))
    add(P("Let *q*_{0} denote the initial user question, which produces the response *a*_{0}. In a "
          "direct interaction, the evolving context can be written as"))
    add(EQ('*C*_{1} = {*q*_{0}, *a*_{0}}', 1))
    add(P_noindent("A follow-up question *q*_{1} yields an answer *a*_{1}, extending the context:"))
    add(EQ('*C*_i{t}_{+1} = *C*_i{t} ∪ {*q*_i{t}, *a*_i{t}}', 2))
    add(P_noindent("The difficulty arises when the full set of latent questions relevant to the "
                   "decision, denoted *Q*^{*}, is much larger than the set of questions *Q*_i{u} that "
                   "the user actually asks:"))
    add(EQ('*Q*_i{u} ⊂ *Q*^{∗}', 3))
    add(P_noindent("When the missing subset *Q*^{*} ∖ *Q*_i{u} contains an essential question, such "
                   "as “Does an existing solution already meet the objective?”, the model may "
                   "produce technically sound output while omitting information that would change "
                   "the decision. We call this the *prompt-echo effect*: the explicit question "
                   "determines the shape of the response (Fig. 1). The model may appear to "
                   "“know” the crucial information once prompted, even though it went "
                   "unexpressed throughout a long interaction."))
    ext(FIGURE(doc, "fig1_prompt_echo.png", 1,
               "The prompt-echo effect and resulting tunnel vision. Explicit user questions "
               "illuminate only a limited portion of the decision-relevant information space; "
               "unasked branches (right) never enter the active reasoning trajectory."))
    add(H2("C", "Amplification in Autonomous Systems"))
    add(P("A single incomplete answer may have limited consequences. An autonomous system that "
          "repeats this failure across many steps, whether invoking tools, modifying code, making "
          "purchases, or acting within clinical workflows, can cause serious harm. As models become "
          "more locally competent, they can execute flawed high-level plans more efficiently. We "
          "summarize this failure state as"))
    add(el('<w:p><w:pPr><w:spacing w:before="60" w:after="60"/><w:jc w:val="center"/></w:pPr>'
           + run("high local competence + incomplete global framing", italic=True)
           + "<w:r><w:br/></w:r>" + run("= accelerated tunnel vision.", italic=True) + "</w:p>"))

    # III. Related work ----------------------------------------------------
    add(H1("III", "Related Work"))
    add(H2("A", "Structured and Decomposed Reasoning"))
    add(P("Several lines of research aim to move beyond a single linear reasoning path. "
          "Chain-of-thought prompting [11] and its zero-shot variant [12] showed that eliciting "
          "intermediate reasoning steps improves multistep problem solving. Self-Ask [3] trains or "
          "prompts models to generate follow-up questions and intermediate answers before a final "
          "answer, showing that models may possess the knowledge a task requires yet fail to "
          "compose it; explicitly generating follow-up questions narrows this compositionality "
          "gap."))
    add(P("Tree of Thoughts (ToT) [4] explores and evaluates multiple reasoning paths instead of "
          "committing to a single token-level trajectory. Least-to-Most prompting [5] decomposes "
          "complex problems into simpler subproblems. Plan-and-Solve prompting [6] first devises a "
          "plan and then executes its steps, improving on unstructured zero-shot chain-of-thought "
          "reasoning. Skeleton-of-Thought (SoT) [7] generates a high-level outline before filling "
          "in details, analogous to how people write. Quiet-STaR [8] shows that internal "
          "reasoning-like computation can improve prediction beyond simple token continuation."))
    add(H2("B", "Self-Critique, Verification, and Oversight"))
    add(P("A complementary line of work checks model output after generation. Self-Refine [13] "
          "iteratively critiques and revises a model's own output; Chain-of-Verification [14] "
          "plans and answers verification questions to reduce hallucination; multiagent debate "
          "[15] uses several model instances to challenge one another; and Constitutional AI [16] "
          "uses model-generated critiques against explicit principles to shape behavior. These "
          "approaches motivate the supervisory lobe of Section VIII, but they generally verify "
          "the answer to the question that was asked rather than asking whether the right "
          "question was asked at all."))
    add(H2("C", "Positioning of This Work"))
    add(P("These methods support our argument but do not address it fully. Most operate at "
          "inference time or focus on reasoning toward an answer for the task as posed. Our "
          "objective is to predict the latent question space itself and to construct a "
          "comprehensive context from it. The goal shifts from “solving the problem with more "
          "steps” to “understanding the problem as a whole before deciding on a course "
          "of action.”"))

    # IV. Prompt-space prediction ------------------------------------------
    add(H1("IV", "Proposed Training Objective I: Prompt-Space Prediction"))
    add(H2("A", "Core Concept"))
    add(P("In conventional supervised instruction tuning, given user intent *I* and initial query "
          "*q*_{0}, the model learns"))
    add(EQ('*P*(*a*_{0} | *I*, *q*_{0})', 4))
    add(P_noindent("We propose adding an auxiliary target that captures a structured set of latent "
                   "questions:"))
    add(EQ('*P*(*Q*_i{L} | *I*, *q*_{0})', 5))
    add(P_noindent("where *Q*_i{L} = {*q*_{1}, *q*_{2}, …, *q*_i{k}} denotes questions that the "
                   "user did not state but that are needed for a well-informed response. Relevance "
                   "is what matters: a useful latent question may differ substantially from the "
                   "user's wording yet be central to the user's goal."))
    add(P("For example, a request such as “Help me build X” may raise the following "
          "latent questions:"))
    for t in ["Does X already exist?",
              "Are there well-established open-source alternatives?",
              "Would a commercial or hosted solution be cheaper than building in-house?",
              "Which constraints might invalidate the proposed design?",
              "Which unforeseen dependencies could become bottlenecks?",
              "Which alternative interpretations better capture the user's actual intent?",
              "What evidence would justify revisiting the decision to build?"]:
        add(ITEM("•", t))
    add(P("The training objective should therefore reward identifying relevant neighboring "
          "questions, not merely generating paraphrases of the original one."))
    add(H2("B", "Expanding Context Before Response Generation"))
    add(P("The predicted question space need not be shown to the user. The questions can instead "
          "serve as latent or explicit intermediate variables that broaden the context:"))
    add(EQ('*C*^{+} = *C*_{0} ∪ {*q*_i{i}, *â*_i{i} : *q*_i{i} ∈ *Q*_i{L}}', 6))
    add(P_noindent("where â_i{i} denotes an internally generated or retrieved answer to latent "
                   "question *q*_i{i}. The response is then generated as"))
    add(EQ('*P*(*a*_{0} | *C*^{+})', 7))
    add(P_noindent("so that conditioning extends beyond the narrow initial context *C*_{0} "
                   "(Fig. 2). This reframes the task: the system becomes responsible not only for "
                   "answering the explicit question but also for anticipating which hidden "
                   "questions must be resolved to answer it adequately."))
    ext(FIGURE(doc, "fig2_prompt_space.png", 2,
               "Prompt-space prediction. (a) A continuation-centric model maps the prompt directly "
               "to a detailed answer. (b) The proposed model first predicts a structured set of "
               "latent questions, uses their answers to build an expanded context, and conditions "
               "the response on that context."))
    add(H2("C", "Training Data Construction"))
    add(P("Supervision for prompt-space prediction can be derived from several sources:"))
    for i, t in enumerate([
            "interactions with human experts who record the questions they consider before making "
            "a decision;",
            "long dialogues that reveal which questions arise after the initial difficulties;",
            "counterfactual regret analysis, in which cases where late-discovered information "
            "changed the optimal decision are recast as training examples of the form “Which "
            "earlier question would have revealed this fact?”;",
            "hierarchical structures extracted from textbooks, technical manuals, design "
            "documents, and diagnostic protocols; and",
            "self-generated candidate questions that are subsequently validated."], 1):
        add(ITEM(f"{i})", t))
    add(P("Late-regret cases, in which users realize that substantial effort could have been "
          "avoided by first checking whether an alternative existed or whether a relevant "
          "constraint had been considered, may be especially valuable training data."))
    add(H2("D", "Loss Function"))
    add(P("We express a composite objective as"))
    add(EQ('ℒ = ℒ_{token} + *λ*_i{q}ℒ_{question} + *λ*_i{c}ℒ_{coverage} + *λ*_i{d}ℒ_{decision}', 8))
    add(P_noindent("where ℒ_{token} preserves standard language-modeling ability, "
                   "ℒ_{question} encourages prediction of relevant latent questions, "
                   "ℒ_{coverage} penalizes neglect of essential aspects of the problem, and "
                   "ℒ_{decision} rewards questions whose answers improve the decision. The "
                   "coefficients *λ*_i{q}, *λ*_i{c}, and *λ*_i{d} weight the auxiliary "
                   "terms. Not every term needs to be differentiable; preference optimization or "
                   "reinforcement learning can be used to optimize the coverage and decision "
                   "terms."))

    # V. Coarse-to-fine ----------------------------------------------------
    add(H1("V", "Proposed Training Objective II: Coarse-to-Fine Context Construction"))
    add(H2("A", "Rationale"))
    add(P("Unbounded question expansion is computationally impractical: a model cannot consider "
          "every conceivable question across a large domain. A hierarchical objective is "
          "therefore needed."))
    add(P("People typically manage this complexity by first forming a broad representation. "
          "When sketching a house, one draws the outline before the interior details (Fig. 3). "
          "When writing a paper, one sets the overall structure before composing individual "
          "sentences. In clinical diagnosis, broad syndromic categories are usually identified "
          "before individual hypotheses are examined."))
    ext(FIGURE(doc, "fig3_coarse_to_fine.png", 3,
               "Coarse-to-fine context construction follows an outline-first process: the global "
               "structure is established before progressively finer details are added."))
    add(P("We advocate the following coarse-to-fine procedure:"))
    for i, t in enumerate([
            "*Broad outline*: infer the key dimensions of the user's intent.",
            "*Major branch expansion*: identify the main alternatives, constraints, risks, "
            "dependencies, and objectives.",
            "*Decision-relevance selection*: determine which branches could change the current "
            "plan.",
            "*Targeted refinement*: spend detailed computation or retrieval only on the branches "
            "that require it.",
            "*Action synthesis*: produce the final response or plan from the expanded "
            "representation."], 1):
        add(ITEM(f"{i})", t))
    add(H2("B", "Explicit Coverage Metric"))
    add(P("What matters here is not the depth of the answer but whether the relevant decision "
          "dimensions are included. Let *D* = {*d*_{1}, …, *d*_i{m}} be the significant "
          "problem dimensions, and let *w*_i{i} be the estimated cost of omitting dimension "
          "*d*_i{i}. A coverage score for context *C* can be defined as"))
    add(EQ('Coverage(*C*) = Σ_i{i} *w*_i{i} 1(*d*_i{i} ∈ *C*) / Σ_i{i} *w*_i{i}', 9))
    add(P_noindent("where 1(\u00b7) is the indicator function. Training can then favor high "
                   "coverage under a computational budget, shifting the objective from "
                   "“consider everything” to “cover the high-consequence regions "
                   "first.”"))
    add(H2("C", "Relation to Existing Decomposition Techniques"))
    add(P("Plan-and-Solve [6], Least-to-Most [5], ToT [4], and SoT [7] show that decomposing or "
          "structuring reasoning before detailed generation can improve performance. Our framework "
          "differs in *what* is decomposed. Existing methods decompose the task as stated, "
          "whereas prompt-space prediction first asks whether the stated task correctly captures "
          "the underlying problem."))

    # VI. Launchpad --------------------------------------------------------
    add(H1("VI", "The Launchpad Hypothesis: From Prompt Dependence to Autonomy"))
    add(P("Our central hypothesis is that dependence on the user to define the prompt space may "
          "be a major bottleneck to genuine autonomy. Current systems show impressive "
          "capabilities, including code generation, information retrieval, tool use, planning, "
          "multimodal perception, and long-context processing, but much of this activity still "
          "depends on what the user asks."))
    add(P("If a model learns to perform this expansion itself, its autonomy increases "
          "substantially, because it no longer relies on the user to delimit the entire search "
          "space. This leads to the *launchpad hypothesis* (Fig. 4): removing the Open Sesame "
          "barrier could convert latent capability into substantially greater reasoning autonomy. "
          "This is a hypothesis that requires empirical validation."))
    ext(FIGURE(doc, "fig4_launchpad.png", 4,
               "The launchpad hypothesis: reducing dependence on user-supplied prompts may convert "
               "latent model capability into substantially greater reasoning autonomy, at the price "
               "of a larger cost of error."))

    # VII. Medical AI ------------------------------------------------------
    add(H1("VII", "Medical AI as a Testbed"))
    add(H2("A", "Why Medicine Exposes the Limitation"))
    add(P("Clinical decision-making carries substantial risk from overlooked factors [9]. An "
          "answer that is internally consistent can still be dangerous in the wider clinical "
          "picture. For a dosing request, for example, a correct calculation does not guarantee a "
          "safe outcome. Critical latent questions include renal function, hepatic function, the "
          "validity of the recorded weight, pregnancy status, allergies, drug interactions, "
          "contraindications, the reliability of the working diagnosis, recent changes in the "
          "patient's condition, signs of infection or sepsis, and missing laboratory results."))
    add(P("A system that relies on the clinician to enumerate every relevant branch therefore "
          "lacks the proactivity that high-stakes situations require."))
    add(H2("B", "Safety Implications of Expanded Autonomy"))
    add(P("Mechanisms that reduce the risk of omission can simultaneously increase the risk of "
          "misinformation. As models generate their own questions and broaden their reasoning, "
          "erroneous assumptions can propagate. A 2026 benchmark of LLM-based agent systems on "
          "clinical decision tasks found that, despite access to tools such as web browsing and "
          "code execution, the agents achieved only modest accuracy gains over baseline LLMs, "
          "while consuming substantially more computation and exhibiting persistent "
          "hallucinations [10]."))
    add(P("The safety challenge is therefore asymmetric: the better a model becomes at "
          "expanding its own problem space, the less acceptable it is for that same process to be "
          "the sole arbiter of what is true."))

    # VIII. Dual-lobe ------------------------------------------------------
    add(H1("VIII", "Dual-Lobe Supervision as a Safety Architecture"))
    add(H2("A", "Functional Separation"))
    add(P("We propose a dual-lobe architecture with distinct functional objectives (Fig. 5)."))
    add(P("*Lobe A (generator/executor)* expands the prompt space, builds a coarse-to-fine "
          "representation, reasons through the task, uses tools and external resources, and "
          "produces actionable output."))
    add(P("*Lobe B (independent supervisor)* actively searches for critical omitted branches, "
          "verifies information against external sources, checks that tool use matches what was "
          "reported, detects shifts in user intent, detects fixation on incomplete solutions, "
          "explicitly checks for existing alternatives, enforces domain-specific safety "
          "standards, and decides when to escalate to human oversight."))
    ext(FIGURE(doc, "fig5_dual_lobe.png", 5,
               "Dual-lobe supervision. Lobe A produces a candidate answer from the expanded "
               "context. Lobe B receives the user's intent and the candidate, but uses its own "
               "retrieval, rule sets, and thresholds; it approves the answer, returns it for "
               "revision, or escalates unresolved cases to a human."))
    add(H2("B", "Addressing Solution Fixation"))
    add(P("A prominent risk in long projects is contextual momentum. Over time, both the user and "
          "the model may come to treat one solution path as the goal itself and overlook other "
          "viable options. The supervisory lobe should therefore perform periodic reassessments, "
          "asking questions such as “If we ignored past decisions and considered only the "
          "user's underlying goal, would we still pursue the current approach?” This "
          "counterfactual check directly targets cases in which a system invests substantial "
          "effort in a solution that an early check for existing tools would have made "
          "unnecessary."))
    add(H2("C", "Independence Rather Than Duplication"))
    add(P("The mere presence of a second model does not guarantee effective supervision. If both "
          "lobes share the same assumptions, training biases, or reasoning patterns, they may "
          "reinforce each other's errors. Effective supervision requires a degree of independence, "
          "which can be obtained through different system objectives, different context, "
          "deliberately adversarial questioning, independent information retrieval, different "
          "model families, domain-specific rule sets, or calibrated disagreement thresholds."))
    add(P("The goal is not two parallel voices but two differentiated cognitive roles."))

    # IX. Evaluation -------------------------------------------------------
    add(H1("IX", "Evaluation Framework"))
    add(P("The proposal can be evaluated without training a large model from scratch. Controlled "
          "studies can compare conventional agents, inference-time self-questioning agents, and "
          "systems specifically trained for prompt-space expansion."))
    add(H2("A", "Benchmark 1: Existing-Solution Discovery"))
    add(P("Construct software tasks in which the system is technically able to build a solution, "
          "a mature existing option already satisfies the requirements, and the user never asks "
          "whether such an option exists. Metrics include the time to identify the existing "
          "solution, the number of tool calls before discovery, the total tokens and computation "
          "consumed, the amount of unnecessary implementation work, and the frequency with which "
          "the system raises the alternative unprompted."))
    add(H2("B", "Benchmark 2: Hidden-Constraint Reversal"))
    add(P("Construct scenarios in which omitting a single inferable constraint changes the correct "
          "decision. Metrics include recall of latent questions, accuracy of decision reversals, "
          "the rate of false-positive interruptions, and coverage under a strict computational "
          "budget."))
    add(H2("C", "Benchmark 3: Clinical Omission Pressure Test"))
    add(P("Construct scenarios in which the explicit request invites a narrow answer, but a "
          "critical adjacent factor should modify or override it. Examples include dose "
          "calculation in renal impairment, anticoagulation questions during active bleeding, "
          "fever management in suspected sepsis, and drug selection under pregnancy "
          "contraindications. Metrics include recall of essential latent questions, the rate of "
          "harmful recommendations, calibration, the supervisory lobe's detection rate, and the "
          "rate of unnecessary escalations."))
    add(H2("D", "Benchmark 4: Context Momentum and Solution Fixation"))
    add(P("Lead the system through an extended trajectory in which it invests substantial effort "
          "in a suboptimal path, then introduce evidence favoring a better alternative without "
          "explicitly asking the system to reconsider. Metrics assess whether the system "
          "acknowledges the new evidence, re-evaluates prior commitments, recommends stopping or "
          "switching, and preserves user agency while proposing alternatives."))

    # X. Hypotheses --------------------------------------------------------
    add(H1("X", "Research Questions and Testable Hypotheses"))
    add(P("The proposal raises the following hypotheses, each of which can be tested "
          "empirically with the benchmarks of Section IX."))
    for tag, t in [
            ("H1", "Models trained for prompt-space prediction will surface relevant omitted "
                   "questions more often than standard instruction-tuned models under comparable "
                   "inference budgets."),
            ("H2", "Prompt-space prediction will reduce wasted execution in build-versus-buy and "
                   "existing-solution scenarios."),
            ("H3", "Coarse-to-fine context construction will improve coverage of high-impact "
                   "branches at lower cost than unrestricted self-question generation."),
            ("H4", "Combining the two objectives will improve robustness to ambiguous prompts, "
                   "but may generate irrelevant questions unless training emphasizes decision "
                   "relevance."),
            ("H5", "Dual-lobe supervision will reduce critical omissions and the propagation of "
                   "hallucinated content relative to a single expanded-reasoning model, provided "
                   "the supervisor has independent objectives or information."),
            ("H6", "In clinical evaluations, gains will come mainly from improved recall of "
                   "essential latent questions rather than from improved arithmetic or factual "
                   "recall.")]:
        add(ITEM(tag + ":", t))

    # XI. Limitations ------------------------------------------------------
    add(H1("XI", "Limitations and Risks"))
    add(P("Prompt-space expansion carries several inherent risks:"))
    for t in ["excessive expansion may increase latency and computational cost;",
              "overly proactive systems may frustrate users by raising irrelevant alternatives or "
              "needlessly challenging clear instructions;",
              "latent question generation may invent risks or constraints that do not exist;",
              "hierarchical outlines may prematurely impose an inappropriate conceptual framing on "
              "a problem; and",
              "dual-lobe architectures may create unwarranted confidence if both lobes share the "
              "same errors."]:
        add(ITEM("•", t))
    add(P("The goal is therefore not maximal proactivity but *calibrated* proactivity: expanding "
          "the context when the expected cost of an omission exceeds the cost of the expansion. "
          "An effective system should adapt its expansion to the estimated significance, "
          "reversibility, uncertainty, and expected error cost of each task. Finally, this paper "
          "presents a conceptual framework and evaluation protocol; the hypotheses above remain to "
          "be validated empirically."))

    # XII. Discussion ------------------------------------------------------
    add(H1("XII", "Discussion"))
    add(P("The Open Sesame barrier exposes a gap between what users expect of AI systems and what "
          "current interaction models deliver. Users expect a collaborator that recognizes the "
          "questions they have not asked, yet these systems remain largely driven by explicit "
          "instructions."))
    add(P("As tool use and autonomy increase, this gap becomes more visible. Longer "
          "chain-of-thought output and larger context windows do not address the root problem: a "
          "model can reason deeply within a poorly framed problem. The more important question is "
          "how the problem space is constructed in the first place. Prompt-space prediction "
          "addresses this directly by training the model to infer the latent questions that "
          "follow from a given intent. Coarse-to-fine construction controls combinatorial growth "
          "by ensuring broad coverage before refinement. Dual-lobe supervision guards against the "
          "additional risk that greater autonomy introduces."))
    add(P("Medicine is an especially demanding proving ground, because relevant information is "
          "spread across many interacting variables and the consequences of omission can be "
          "severe. A clinically competent AI system should not assume that a physician will "
          "explicitly ask about renal function, pregnancy, drug interactions, deterioration, or "
          "sepsis before those factors are considered. At the same time, a system that generates "
          "its own questions needs robust oversight, because poorly chosen branches can introduce "
          "substantial risk. The guiding principle is that gains in cognitive capability must be "
          "matched by a correspondingly independent supervisory framework."))

    # XIII. Conclusion -----------------------------------------------------
    add(H1("XIII", "Conclusion"))
    add(P("Current LLMs are highly capable, yet they remain behaviorally dependent on user input "
          "to illuminate critical parts of the problem space. We characterized this dependence as "
          "the Open Sesame barrier, which produces tunnel vision: the model can reason deeply along "
          "an activated trajectory while failing to surface information that would contradict it."))
    add(P("We advocate intervening at the level of training. Models should learn not only to "
          "predict continuations but also to anticipate the latent question space implied by the "
          "user's intent, and to build a coarse-to-fine representation before committing to "
          "detailed actions or answers. Existing methods, including Self-Ask, Tree of Thoughts, "
          "Plan-and-Solve, Least-to-Most, Skeleton-of-Thought, and latent-rationale training, show "
          "that deliberate evaluation and decomposition improve reasoning; our proposal focuses on "
          "surfacing the questions that must be asked before a task can be considered complete."))
    add(P("If successful, this approach could substantially increase model autonomy, which in "
          "turn makes independent oversight necessary. We therefore proposed dual-lobe supervision, "
          "in which a generator/executor performs the reasoning while an independent supervisor "
          "identifies omissions, verifies factual claims, detects shifts in user intent, counters "
          "solution fixation, and enforces domain-specific safety standards."))
    add(P("The practical motivation is clear: in software engineering, a missed question wastes "
          "time and resources; in medicine, it can endanger patients. The long-term goal is an AI "
          "system that not only answers the questions it is asked but also recognizes which "
          "questions should follow, including those the user has not yet realized must be "
          "addressed."))

    # References -----------------------------------------------------------
    add(REF_HEAD())
    for r in REFS:
        add(REF(r))
    return B


def bio():
    txt = (f"{AUTHOR} received the [[degree]] degree in [[field]] from [[university]], "
           "[[city, country]], in [[year]]. [[Current position, department, institution, city, "
           "country.]] [[His/Her/Their]] research interests include large language models, AI "
           "safety, and clinical decision support.")
    first, rest = txt.split(" received", 1)
    return el('<w:p><w:pPr><w:spacing w:before="360"/><w:jc w:val="both"/></w:pPr>'
              f'{run(first, bold=True)}{runs(" received" + rest)}</w:p>')


# ------------------------------ assembly ----------------------------------
def set_runs(p_elm, xml_runs):
    for r in list(p_elm):
        if r.tag != f"{{{W}}}pPr":
            p_elm.remove(r)
    for r in etree.fromstring(f"<root {NS}>{xml_runs}</root>"):
        p_elm.append(r)


def main():
    doc = docx.Document(str(TEMPLATE))
    paras = doc.paragraphs

    # Title
    set_runs(paras[0]._p, run(TITLE))
    # Authors (single author; IEEE membership grade goes after the comma if applicable)
    set_runs(paras[2]._p, run(AUTHOR, size=22))
    # Abstract
    set_runs(paras[6]._p, run("Abstract", bold=True, italic=True, size=18)
             + run("—" + ABSTRACT, bold=True, size=18))
    # Index terms
    set_runs(paras[8]._p, run("Index Terms", bold=True, italic=True, size=18)
             + run("—" + INDEX_TERMS + ".", bold=True, size=18))

    # Replace every template instruction paragraph after the index terms

    anchor = paras[9]._p
    for p in paras[10:]:
        p._p.getparent().remove(p._p)
    new = body(doc) + [bio()]
    for e in reversed(new):
        anchor.addnext(e)

    # First-page footnote (funding, corresponding author, affiliation)
    fn_part = next(r.target_part for r in doc.part.rels.values() if r.reltype.endswith("/footnotes"))
    fn_xml = etree.fromstring(fn_part.blob)
    note = fn_xml.xpath('//w:footnote[@w:id="1"]', namespaces={"w": W})[0]
    ps = note.xpath("w:p", namespaces={"w": W})
    template_p = copy.deepcopy(ps[0])
    for p in ps:
        note.remove(p)
    lines = [
        "[[Funding statement, e.g., \u201cThis work received no external funding.\u201d]] "
        "(Corresponding author: " + AUTHOR + ".)",
        "A. Alsawy is with [[department]], [[institution]], [[city, postal code, country]] "
        "(e-mail: [[e-mail address]]).",
    ]
    for i, line in enumerate(lines):
        p = copy.deepcopy(template_p)
        for r in p.xpath("w:r|w:hyperlink|w:proofErr|w:bookmarkStart|w:bookmarkEnd",
                         namespaces={"w": W}):
            is_mark = bool(r.xpath(".//w:footnoteRef|.//w:sym", namespaces={"w": W}))
            if not (i == 0 and is_mark):
                p.remove(r)
        for r in etree.fromstring(f"<root {NS}>{runs(line, size=16)}</root>"):
            p.append(r)
        note.append(p)
    fn_part._blob = etree.tostring(fn_xml, xml_declaration=True, encoding="UTF-8", standalone=True)

    # Running header

    hdr = doc.sections[0].header
    for p in hdr.paragraphs:
        if "REPLACE THIS LINE" in p.text:
            ts = p._p.xpath(".//w:t")
            ts[0].text = f"> {RUNNING_HEAD} <"
            for t in ts[1:]:
                t.text = ""

    # Document properties
    cp = doc.core_properties
    cp.title = TITLE
    cp.author = AUTHOR
    cp.subject = "Submitted to IEEE Transactions on Neural Networks and Learning Systems"
    cp.keywords = INDEX_TERMS
    cp.comments = ""

    doc.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
